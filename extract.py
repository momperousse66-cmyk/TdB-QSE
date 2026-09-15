import openpyxl,json,datetime,warnings
warnings.simplefilter('ignore')
p=r'C:/Users/mompe/Downloads/Lille_TDB QSE.xlsm'
w=openpyxl.load_workbook(p,read_only=True,data_only=True)
f=openpyxl.load_workbook(p,read_only=True,data_only=False)
def clean(v):
 if isinstance(v,(datetime.datetime,datetime.date,datetime.time)): return v.isoformat()
 return v
data={s.title:[[clean(c.value) for c in row] for row in s] for s in w}
forms={s.title:{c.coordinate:(c.value if isinstance(c.value,str) else c.value.text) for row in s for c in row if c.data_type=='f'} for s in f}
json.dump({'data':data,'formulas':forms},open('source.json','w',encoding='utf8'),ensure_ascii=False,indent=2)
for name in ['Indicateurs','Suivi Habilitations','Tableau suivi RC ','Données RC','Donné Restrictions','Restrictions']:
 print('\nSHEET',name)
 for i,row in enumerate(data[name]):
  vals=['%s:%s'%(openpyxl.utils.get_column_letter(j+1),v) for j,v in enumerate(row) if v is not None]
  if vals and (name=='Indicateurs' or i<14):print(i+1,' '.join(vals))
print('\nPDP FORMULAS',forms['Suivi PdP'])
print('\nQUALITY FORMULAS', {k:v for k,v in forms['Indice qualité et nb contrôle'].items() if k in ['I2','I5','O5']})
print('\nERRORS',[(n,sum(isinstance(v,str) and v.startswith(('#REF!','#DIV/0!','#VALUE!','#N/A')) for row in rows for v in row)) for n,rows in data.items()])

